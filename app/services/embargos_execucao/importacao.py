"""Leitura do relatório "ROBÔ - EMBARGOS À EXECUÇÃO" do L1 e da planilha do legado.

Colunas casadas pelo TÍTULO, nunca pela posição: o operador ajusta o modelo
799 no L1 e a planilha do legado tem layout próprio. No relatório de 11/09/2026
os títulos úteis eram "Id" (id da tarefa), "Data/hora conclusão efetiva"
(= ajuizamento), "NÚMERO DO PROCESO", "NÚMERO DA PASTA", "Vínculos com
processo / Título" (NPJ), "ESCRITÓRIO RESPONSÁVEL", "ADOVOGADO(A) RESPONSÁVEL"
(com o erro de digitação do modelo), "EXECUTANTE", "UF" e "STATUS".
Armadilhas: "CONCLUSÃO PREVISTA" é o PRAZO, não a conclusão; "Compromissos
gerados / Envolvidos / É executante" não é o executante.
"""
from __future__ import annotations

import csv
import io
import logging
from typing import Any, Optional

from sqlalchemy.orm import Session

from app.models.embargos_execucao import EVT_AVISO, EVT_INFO, ORIGEM_RELATORIO, SECAO_ENTRADA
from app.services.embargos_execucao import service
from app.services.embargos_execucao.normaliza import para_data, sem_acento

logger = logging.getLogger(__name__)

# (campo, prioridade) — menor prioridade vence quando dois títulos casam.
def _campo_do_titulo(titulo: Any) -> Optional[tuple[str, int]]:
    n = " ".join(sem_acento(titulo).lower().split())
    if not n:
        return None
    if n in ("id", "id da tarefa", "id tarefa", "codigo da tarefa", "id do compromisso/tarefa"):
        return ("l1_task_id", 0)
    if "compromissos gerados" in n or "prevista" in n or "previsto" in n:
        return None
    if "conclusao efetiva" in n:
        return ("data_ajuizamento", 0)
    if "ajuizamento" in n:
        return ("data_ajuizamento", 1)
    if "distribuicao" in n or "protocolo" in n:
        return ("data_ajuizamento", 2)
    if "titulo" in n or "npj" in n:
        return ("npj", 0)
    if "numero da pasta" in n:
        return ("pasta", 0)
    if n in ("pasta", "proc"):
        return ("pasta", 1)
    if "cnj" in n or "numero do proces" in n or n in ("processo", "numero processo"):
        return ("cnj", 0)
    if "escritorio" in n:
        return ("escritorio", 0)
    if "executante" in n:
        return ("executante_nome", 0)
    if "advogad" in n or "adovogad" in n or "responsavel" in n:
        return ("responsavel_nome", 0)
    if n == "uf":
        return ("uf", 0)
    if n == "status":
        return ("status", 0)
    return None


def mapear_cabecalho(cabecalho: tuple | list) -> dict[str, int]:
    melhor: dict[str, tuple[int, int]] = {}
    for idx, titulo in enumerate(cabecalho or ()):
        achado = _campo_do_titulo(titulo)
        if not achado:
            continue
        campo, prio = achado
        if campo not in melhor or prio < melhor[campo][0]:
            melhor[campo] = (prio, idx)
    return {campo: idx for campo, (_, idx) in melhor.items()}


def _texto(valor: Any) -> Optional[str]:
    if valor is None:
        return None
    s = str(valor).strip()
    return s or None


def linhas_da_tabela(rows: list[tuple]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Acha o cabeçalho (primeiras 20 linhas) e devolve as linhas úteis + contagem."""
    cab_idx, mapa = None, {}
    for i, row in enumerate(rows[:20]):
        m = mapear_cabecalho(row)
        if ("pasta" in m or "cnj" in m) and len(m) >= 2:
            cab_idx, mapa = i, m
            break
    stats = {"cabecalho_encontrado": cab_idx is not None, "colunas": sorted(mapa),
             "linhas_lidas": 0, "status_ignorado": 0, "sem_data": 0, "sem_identificador": 0}
    if cab_idx is None:
        return [], stats

    def val(row, campo):
        i = mapa.get(campo)
        return row[i] if i is not None and i < len(row) else None

    saida: list[dict[str, Any]] = []
    for row in rows[cab_idx + 1:]:
        if not row or not any(c not in (None, "") for c in row):
            continue
        stats["linhas_lidas"] += 1
        status = _texto(val(row, "status"))
        if status and sem_acento(status).strip().lower() != "cumprido":
            stats["status_ignorado"] += 1
            continue
        pasta, cnj = _texto(val(row, "pasta")), _texto(val(row, "cnj"))
        if not pasta and not cnj:
            stats["sem_identificador"] += 1
            continue
        data = para_data(val(row, "data_ajuizamento"))
        if not data:
            stats["sem_data"] += 1
            continue
        task_id = val(row, "l1_task_id")
        try:
            task_id = int(str(task_id).strip()) if task_id not in (None, "") else None
        except ValueError:
            task_id = None
        saida.append({
            "pasta": pasta,
            "cnj": cnj,
            "npj": _texto(val(row, "npj")),
            "data_ajuizamento": data,
            "l1_task_id": task_id,
            "escritorio": _texto(val(row, "escritorio")),
            "responsavel_nome": _texto(val(row, "responsavel_nome")),
            "executante_nome": _texto(val(row, "executante_nome")),
            "uf": _texto(val(row, "uf")),
        })
    return saida, stats


def ler_tabela(conteudo: bytes, nome_arquivo: str) -> list[tuple]:
    nome = (nome_arquivo or "").lower()
    if nome.endswith((".csv", ".txt")):
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                texto = conteudo.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        amostra = (texto or "")[:4096]
        try:
            dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        return [tuple(r) for r in csv.reader(io.StringIO(texto or ""), dialeto)]

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        # Primeira aba que tiver cabeçalho reconhecível.
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            linhas, stats = linhas_da_tabela(rows)
            if stats["cabecalho_encontrado"]:
                return rows
        return list(wb.worksheets[0].iter_rows(values_only=True)) if wb.worksheets else []
    finally:
        wb.close()


def importar_bytes(
    db: Session,
    conteudo: bytes,
    nome_arquivo: str,
    *,
    origem: str,
    user_id: Optional[int] = None,
) -> dict[str, Any]:
    rows = ler_tabela(conteudo, nome_arquivo)
    linhas, stats = linhas_da_tabela(rows)
    if not stats["cabecalho_encontrado"]:
        raise ValueError(
            "Não achei o cabeçalho: a planilha precisa de uma coluna de pasta ou de "
            "número do processo (CNJ) e da data do ajuizamento/conclusão."
        )
    if "data_ajuizamento" not in stats["colunas"]:
        raise ValueError(
            "Falta a coluna da data do ajuizamento (ex.: 'Data/hora conclusão efetiva' "
            "ou 'Data do ajuizamento')."
        )
    antes_do_corte = 0
    corte = None
    if origem == ORIGEM_RELATORIO:
        # O relatório traz o histórico inteiro do subtipo: só entra caso NOVO.
        corte = service.corte_relatorio(definir_se_vazio=True)
        antes_do_corte = sum(1 for l in linhas if l["data_ajuizamento"] < corte)
        linhas = [l for l in linhas if l["data_ajuizamento"] >= corte]
    resultado = service.upsert_execucoes(db, linhas, origem=origem, user_id=user_id)
    resultado.update({k: stats[k] for k in ("linhas_lidas", "status_ignorado", "sem_data", "sem_identificador", "colunas")})
    resultado["antes_do_corte"] = antes_do_corte
    resultado["corte"] = corte.isoformat() if corte else None
    service.registrar_evento(
        db, SECAO_ENTRADA,
        f"Importação ({origem.lower()}) de '{nome_arquivo}': {resultado['novas']} nova(s), "
        f"{resultado['atualizadas']} atualizada(s), {stats['sem_data']} sem data, "
        f"{stats['sem_identificador']} sem pasta/CNJ"
        + (f", {antes_do_corte} anterior(es) ao corte de {corte:%d/%m/%Y} (ignoradas)." if corte else "."),
        nivel=EVT_AVISO if (stats["sem_data"] or stats["sem_identificador"]) else EVT_INFO,
        dados=resultado, user_id=user_id,
    )
    db.commit()
    return resultado
